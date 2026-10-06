"""Lógica de negocio: sin nada de interfaz."""
import shutil
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook

COLUMNAS = ["codigo", "nombre", "categoria", "subcategoria", "unidad",
            "contenido", "precio_compra", "precio_venta", "stock", "stock_minimo"]

# ---------- Catálogo ----------
def categoria_id(conn, nombre, padre_id=None):
    if not nombre:
        return padre_id
    nombre = nombre.strip().title()
    r = conn.execute("SELECT id FROM categorias WHERE nombre=? AND padre_id IS ?",
                     (nombre, padre_id)).fetchone()
    if r:
        return r["id"]
    return conn.execute("INSERT INTO categorias(nombre,padre_id) VALUES(?,?)",
                        (nombre, padre_id)).lastrowid

def listar_categorias(conn):
    """Devuelve [(id, 'Limpieza > Detergente')]"""
    return [(r["id"], r["ruta"]) for r in conn.execute(
        """SELECT c.id, COALESCE(p.nombre||' > ','')||c.nombre AS ruta
           FROM categorias c LEFT JOIN categorias p ON p.id=c.padre_id ORDER BY ruta""")]

def guardar_producto(conn, d):
    with conn:
        cur = conn.execute(
            """INSERT INTO productos(codigo_barras,nombre,categoria_id,unidad_medida,contenido,
               precio_compra,precio_venta,stock_minimo) VALUES(?,?,?,?,?,?,?,?)""",
            (d.get("codigo") or None, d["nombre"], d.get("categoria_id"),
             d.get("unidad", "unidad"), d.get("contenido", 0), d.get("precio_compra", 0),
             d["precio_venta"], d.get("stock_minimo", 5)))
        pid = cur.lastrowid
        if d.get("stock", 0) > 0:
            mover_stock(conn, pid, "entrada", d["stock"], "carga inicial")
    return pid

def buscar_productos(conn, texto, limite=50):
    like = f"%{texto}%"
    return conn.execute(
        """SELECT p.*, COALESCE(c.nombre,'') categoria FROM productos p
           LEFT JOIN categorias c ON c.id=p.categoria_id
           WHERE p.activo=1 AND (p.codigo_barras=? OR p.nombre LIKE ?)
           ORDER BY p.nombre LIMIT ?""", (texto, like, limite)).fetchall()

def listar_productos(conn):
    return conn.execute(
        """SELECT p.*, COALESCE(pc.nombre||' > ','')||COALESCE(c.nombre,'') categoria
           FROM productos p LEFT JOIN categorias c ON c.id=p.categoria_id
           LEFT JOIN categorias pc ON pc.id=c.padre_id
           WHERE p.activo=1 ORDER BY categoria, p.nombre""").fetchall()

# ---------- Inventario ----------
def mover_stock(conn, producto_id, tipo, cantidad, motivo, ref=None):
    """Única vía para cambiar el stock: siempre deja rastro (kardex)."""
    signo = -1 if tipo == "salida" else 1
    conn.execute("UPDATE productos SET stock_actual=stock_actual+? WHERE id=?",
                 (signo * cantidad, producto_id))
    conn.execute("""INSERT INTO movimientos_inventario(producto_id,tipo,cantidad,motivo,referencia_id)
                    VALUES(?,?,?,?,?)""", (producto_id, tipo, cantidad, motivo, ref))

def entrada_stock(conn, producto_id, cantidad, motivo="compra"):
    with conn:
        mover_stock(conn, producto_id, "entrada", cantidad, motivo)

def stock_bajo(conn):
    return conn.execute("""SELECT * FROM productos WHERE activo=1
                           AND stock_actual<=stock_minimo ORDER BY stock_actual""").fetchall()

# ---------- Caja ----------
def turno_abierto(conn):
    return conn.execute("""SELECT *, datetime(apertura,'localtime') apertura_local
                           FROM turnos_caja WHERE estado='abierto'""").fetchone()

def abrir_turno(conn, monto_inicial, cajero="Cajero"):
    if turno_abierto(conn):
        raise ValueError("Ya hay un turno abierto")
    with conn:
        return conn.execute("INSERT INTO turnos_caja(cajero,monto_inicial) VALUES(?,?)",
                            (cajero, monto_inicial)).lastrowid

def registrar_movimiento_caja(conn, tipo, concepto, monto):
    t = turno_abierto(conn)
    if not t:
        raise ValueError("Abre la caja primero")
    with conn:
        conn.execute("INSERT INTO movimientos_caja(turno_id,tipo,concepto,monto) VALUES(?,?,?,?)",
                     (t["id"], tipo, concepto, monto))

def resumen_turno(conn, turno_id):
    t = conn.execute("SELECT * FROM turnos_caja WHERE id=?", (turno_id,)).fetchone()
    por_metodo = {r["metodo"]: r["s"] for r in conn.execute(
        """SELECT p.metodo, SUM(p.monto) s FROM pagos p JOIN ventas v ON v.id=p.venta_id
           WHERE v.turno_id=? AND v.estado='completada' GROUP BY p.metodo""", (turno_id,))}
    mov = {r["tipo"]: r["s"] for r in conn.execute(
        "SELECT tipo, SUM(monto) s FROM movimientos_caja WHERE turno_id=? GROUP BY tipo",
        (turno_id,))}
    efectivo = por_metodo.get("efectivo", 0)
    esperado = t["monto_inicial"] + efectivo + mov.get("ingreso", 0) - mov.get("egreso", 0)
    return {"turno": t, "por_metodo": por_metodo, "ingresos": mov.get("ingreso", 0),
            "egresos": mov.get("egreso", 0), "esperado": esperado,
            "total_ventas": sum(por_metodo.values())}

def cerrar_turno(conn, monto_contado):
    t = turno_abierto(conn)
    if not t:
        raise ValueError("No hay turno abierto")
    with conn:
        conn.execute("""UPDATE turnos_caja SET estado='cerrado', cierre=CURRENT_TIMESTAMP,
                        monto_contado_cierre=? WHERE id=?""", (monto_contado, t["id"]))
    r = resumen_turno(conn, t["id"])
    r["diferencia"] = monto_contado - r["esperado"]
    return r

# ---------- Ventas ----------
def registrar_venta(conn, items, pagos, descuento=0.0, tasa_impuesto=0.0, recibido=None):
    """items: [{producto_id, cantidad, precio}]  pagos: [{metodo, monto}]"""
    t = turno_abierto(conn)
    if not t:
        raise ValueError("Debes abrir la caja antes de vender")
    if not items:
        raise ValueError("El carrito está vacío")
    with conn:  # transacción: todo o nada
        subtotal = sum(i["cantidad"] * i["precio"] for i in items)
        impuesto = round((subtotal - descuento) * tasa_impuesto, 2)
        total = round(subtotal - descuento + impuesto, 2)
        if sum(p["monto"] for p in pagos) + 0.005 < total:
            raise ValueError("Los pagos no cubren el total")
        venta_id = conn.execute(
            "INSERT INTO ventas(turno_id,subtotal,descuento,impuesto,total,recibido) VALUES(?,?,?,?,?,?)",
            (t["id"], subtotal, descuento, impuesto, total, recibido if recibido else total)).lastrowid
        for i in items:
            r = conn.execute("SELECT nombre,stock_actual FROM productos WHERE id=?",
                             (i["producto_id"],)).fetchone()
            if r["stock_actual"] < i["cantidad"]:
                raise ValueError(f"Stock insuficiente: {r['nombre']} (hay {r['stock_actual']})")
            conn.execute("""INSERT INTO venta_detalle(venta_id,producto_id,cantidad,
                            precio_unitario,subtotal) VALUES(?,?,?,?,?)""",
                         (venta_id, i["producto_id"], i["cantidad"], i["precio"],
                          i["cantidad"] * i["precio"]))
            mover_stock(conn, i["producto_id"], "salida", i["cantidad"], "venta", venta_id)
        # el monto registrado es lo que realmente queda en caja (sin el vuelto)
        restante = total
        for p in pagos:
            m = min(p["monto"], restante)
            restante -= m
            conn.execute("INSERT INTO pagos(venta_id,metodo,monto) VALUES(?,?,?)",
                         (venta_id, p["metodo"], m))
    return venta_id, total

def anular_venta(conn, venta_id):
    with conn:
        v = conn.execute("SELECT estado FROM ventas WHERE id=?", (venta_id,)).fetchone()
        if not v or v["estado"] != "completada":
            raise ValueError("Venta no válida para anular")
        for d in conn.execute("SELECT * FROM venta_detalle WHERE venta_id=?", (venta_id,)):
            mover_stock(conn, d["producto_id"], "entrada", d["cantidad"], "anulación", venta_id)
        conn.execute("UPDATE ventas SET estado='anulada' WHERE id=?", (venta_id,))

# ---------- Excel ----------
def crear_plantilla_excel(ruta):
    wb = Workbook(); ws = wb.active; ws.title = "Productos"
    ws.append(COLUMNAS)
    ws.append(["7750001", "Detergente Ariel 500g", "Limpieza", "Detergente", "g", 500, 8, 12, 24, 5])
    ws.append(["7750002", "Atún en lata 170g", "Comida", "Enlatados", "g", 170, 6, 9.5, 48, 10])
    wb.save(ruta)

def importar_excel(conn, ruta, backup_dir=None):
    """Devuelve (importados, errores). Hace backup antes de tocar nada."""
    hacer_backup(conn, backup_dir, prefijo="pre_import")
    ws = load_workbook(ruta, data_only=True).active
    filas = list(ws.iter_rows(values_only=True))
    enc = [str(c).strip().lower() if c else "" for c in filas[0]]
    faltan = {"nombre", "precio_venta"} - set(enc)
    if faltan:
        return 0, [f"Faltan columnas obligatorias: {', '.join(faltan)}"]
    ok, errores = 0, []
    for n, fila in enumerate(filas[1:], start=2):
        d = dict(zip(enc, fila))
        try:
            if not d.get("nombre"):
                continue
            num = lambda k, dflt=0: float(d.get(k) or dflt)
            with conn:
                padre = categoria_id(conn, d.get("categoria"))
                cat = categoria_id(conn, d.get("subcategoria"), padre) if d.get("subcategoria") else padre
                cod = str(d["codigo"]).strip() if d.get("codigo") else None
                existe = conn.execute("SELECT id FROM productos WHERE codigo_barras=?",
                                      (cod,)).fetchone() if cod else None
                datos = (str(d["nombre"]).strip(), cat, d.get("unidad") or "unidad",
                         num("contenido"), num("precio_compra"), num("precio_venta"),
                         num("stock_minimo", 5))
                if existe:  # actualiza precios; el stock se suma como entrada
                    conn.execute("""UPDATE productos SET nombre=?,categoria_id=?,unidad_medida=?,
                        contenido=?,precio_compra=?,precio_venta=?,stock_minimo=? WHERE id=?""",
                        datos + (existe["id"],))
                    pid = existe["id"]
                else:
                    pid = conn.execute("""INSERT INTO productos(nombre,categoria_id,unidad_medida,
                        contenido,precio_compra,precio_venta,stock_minimo,codigo_barras)
                        VALUES(?,?,?,?,?,?,?,?)""", datos + (cod,)).lastrowid
                if num("stock") > 0:
                    mover_stock(conn, pid, "entrada", num("stock"), "carga Excel")
            ok += 1
        except Exception as e:
            errores.append(f"Fila {n}: {e}")
    return ok, errores

# ---------- Dashboard ----------
def resumen_dashboard(conn):
    hoy = conn.execute("""SELECT COALESCE(SUM(total),0) total, COUNT(*) n FROM ventas
                          WHERE date(fecha,'localtime')=date('now','localtime')
                          AND estado='completada'""").fetchone()
    top = conn.execute("""SELECT p.nombre, SUM(d.cantidad) c FROM venta_detalle d
        JOIN ventas v ON v.id=d.venta_id JOIN productos p ON p.id=d.producto_id
        WHERE v.estado='completada' AND date(v.fecha,'localtime')=date('now','localtime')
        GROUP BY p.id ORDER BY c DESC LIMIT 5""").fetchall()
    return {"ventas_hoy": hoy["total"], "tickets": hoy["n"],
            "ticket_prom": hoy["total"] / hoy["n"] if hoy["n"] else 0,
            "stock_bajo": len(stock_bajo(conn)), "top": top,
            "caja_abierta": turno_abierto(conn) is not None}

# ---------- Backup ----------
def hacer_backup(conn, destino=None, prefijo="backup", mantener=30):
    import sqlite3
    destino = destino or Path(__file__).parent / "backups"   # siempre junto al programa
    Path(destino).mkdir(parents=True, exist_ok=True)
    ruta = Path(destino) / f"{prefijo}_{datetime.now():%Y%m%d_%H%M%S}.db"
    dest = sqlite3.connect(ruta)
    with dest:
        conn.backup(dest)   # copia segura aunque la base esté en uso
    dest.close()
    viejos = sorted(Path(destino).glob(f"{prefijo}_*.db"))[:-mantener]
    for v in viejos:
        v.unlink()
    return ruta


# ---------- Ajustes (tema, negocio, impuesto) ----------
def get_ajuste(conn, clave, defecto=""):
    r = conn.execute("SELECT valor FROM ajustes WHERE clave=?", (clave,)).fetchone()
    return r["valor"] if r else defecto

def set_ajuste(conn, clave, valor):
    with conn:
        conn.execute("""INSERT INTO ajustes(clave,valor) VALUES(?,?)
                        ON CONFLICT(clave) DO UPDATE SET valor=excluded.valor""", (clave, str(valor)))

# ---------- Editar / eliminar productos ----------
def obtener_producto(conn, pid):
    return conn.execute(
        """SELECT p.*, COALESCE(pc.nombre||' > ','')||COALESCE(c.nombre,'') categoria
           FROM productos p LEFT JOIN categorias c ON c.id=p.categoria_id
           LEFT JOIN categorias pc ON pc.id=c.padre_id WHERE p.id=?""", (pid,)).fetchone()

def actualizar_producto(conn, pid, d):
    with conn:
        conn.execute("""UPDATE productos SET codigo_barras=?,nombre=?,categoria_id=?,unidad_medida=?,
            contenido=?,precio_compra=?,precio_venta=?,stock_minimo=? WHERE id=?""",
            (d.get("codigo") or None, d["nombre"], d.get("categoria_id"), d.get("unidad", "unidad"),
             d.get("contenido", 0), d.get("precio_compra", 0), d["precio_venta"],
             d.get("stock_minimo", 5), pid))
        if "stock" in d:   # si cambió el stock a mano, queda registrado como ajuste
            actual = conn.execute("SELECT stock_actual FROM productos WHERE id=?", (pid,)).fetchone()[0]
            dif = d["stock"] - actual
            if abs(dif) > 1e-9:
                mover_stock(conn, pid, "entrada" if dif > 0 else "salida", abs(dif), "ajuste manual")

def desactivar_producto(conn, pid):
    """'Eliminar' sin perder el historial: el producto deja de aparecer y de venderse."""
    with conn:
        conn.execute("""UPDATE productos SET activo=0,
            codigo_barras=CASE WHEN codigo_barras IS NULL THEN NULL
                               ELSE codigo_barras||'#baja'||id END WHERE id=?""", (pid,))

# ---------- Ventas y notas de compra ----------
def obtener_venta(conn, venta_id):
    v = conn.execute("""SELECT v.*, datetime(v.fecha,'localtime') fecha_local,
        COALESCE(t.cajero,'') cajero FROM ventas v
        LEFT JOIN turnos_caja t ON t.id=v.turno_id WHERE v.id=?""", (venta_id,)).fetchone()
    items = conn.execute("""SELECT d.*, p.nombre, p.unidad_medida FROM venta_detalle d
        JOIN productos p ON p.id=d.producto_id WHERE d.venta_id=? ORDER BY d.id""", (venta_id,)).fetchall()
    pagos = conn.execute("SELECT * FROM pagos WHERE venta_id=?", (venta_id,)).fetchall()
    return {"venta": v, "items": items, "pagos": pagos}

def listar_ventas(conn, limite=200):
    return conn.execute("""SELECT v.id, datetime(v.fecha,'localtime') fecha_local,
        COALESCE(t.cajero,'') cajero, v.total, v.estado,
        (SELECT group_concat(DISTINCT metodo) FROM pagos WHERE venta_id=v.id) metodos
        FROM ventas v LEFT JOIN turnos_caja t ON t.id=v.turno_id
        ORDER BY v.id DESC LIMIT ?""", (limite,)).fetchall()

def listar_movimientos_caja(conn, turno_id):
    return conn.execute("""SELECT tipo, concepto, monto,
        strftime('%H:%M', datetime(fecha,'localtime')) hora
        FROM movimientos_caja WHERE turno_id=? ORDER BY id DESC""", (turno_id,)).fetchall()

def ventas_ultimos_dias(conn, n=7):
    from datetime import date, timedelta
    filas = {r["d"]: r["t"] for r in conn.execute(
        """SELECT date(fecha,'localtime') d, SUM(total) t FROM ventas
           WHERE estado='completada' AND date(fecha,'localtime')>=date('now','localtime',?)
           GROUP BY d""", (f"-{n-1} days",))}
    hoy = date.today()
    dias = [hoy - timedelta(days=k) for k in range(n - 1, -1, -1)]
    return [(d, filas.get(d.isoformat(), 0.0)) for d in dias]
